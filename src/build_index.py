"""Static heuristic index: V-Dem democracy scores + HDI (with an HDI regression) -> one composite.

Outputs
  data/processed/index_panel.csv        flat panel (python-computed, mirrors the workbook formulas)
  output/democratic_resilience_index.xlsx  weights + regression + one sheet per case (live formulas)
  output/index_<case>.png / output/index_all_cases.png   plots for the visual check
Run:  python src/build_index.py
"""
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.ticker import MaxNLocator  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROC = ROOT / "data" / "processed"
OUT = ROOT / "output"
PROC.mkdir(parents=True, exist_ok=True)
OUT.mkdir(exist_ok=True)

# ---- Manual, editable heuristics (mirrored in the workbook's Weights sheet) --------------------
WEIGHTS = {  # must sum to 1
    "polyarchy": 0.30,  # V-Dem electoral democracy index
    "libdem": 0.30,     # V-Dem liberal democracy index (constraints on executive, rights)
    "hdi": 0.20,        # development level
    "hdi_adj": 0.20,    # democracy relative to what HDI predicts (regression residual score)
}
STRAIN_THRESHOLD = -0.02  # GDP-per-capita yoy change at/below this = "strain year"
CASES = {  # name: (first year, last year, strongman/regime-change marker year, label)
    "Germany": (1919, 1938, 1933, "Weimar Germany (Hitler appointed 1933)"),
    "Venezuela": (1990, 2023, 1999, "Venezuela (Chavez takes office 1999)"),
    "Hungary": (1990, 2023, 2010, "Hungary (Orban returns 2010)"),
}


def owid(name, col):
    d = pd.read_csv(RAW / f"{name}.csv")
    return d.rename(columns={col: name})[["entity", "year", name]]


def load():
    poly = owid("electoral-democracy-index", "electdem_vdem__estimate_best")
    lib = owid("liberal-democracy-index", "libdem_vdem__estimate_best")
    undp = owid("human-development-index", "hdi__sex_total")
    esc = owid("human-development-index-escosura", "ahdi")
    gdp = owid("gdp-per-capita-maddison-project-database", "gdp_per_capita")
    return poly, lib, undp, esc, gdp


def fit(x, y):
    b, a = np.polyfit(x, y, 1)
    r2 = np.corrcoef(x, y)[0, 1] ** 2
    return a, b, r2, len(x)


def main():
    poly, lib, undp, esc, gdp = load()
    dem = poly.merge(lib, on=["entity", "year"])

    # ---- HDI regressions: liberal democracy ~ HDI, pooled over every country-year ---------------
    hist = dem.merge(esc, on=["entity", "year"]).dropna()  # Escosura (1870-2020, sparse pre-1950)
    modern = dem.merge(undp, on=["entity", "year"]).dropna()  # UNDP (1990-2023, annual)
    reg = {
        "historical": fit(hist["human-development-index-escosura"], hist["liberal-democracy-index"]),
        "modern": fit(modern["human-development-index"], modern["liberal-democracy-index"]),
    }

    panels = []
    for country, (y0, y1, _, _) in CASES.items():
        years = pd.DataFrame({"year": range(y0, y1 + 1)})
        d = years.merge(dem[dem.entity == country], on="year", how="left")
        # HDI: UNDP where it exists (1990+), otherwise Escosura, linearly interpolated between benchmarks
        u = undp[undp.entity == country].set_index("year")["human-development-index"]
        e = esc[esc.entity == country].set_index("year")["human-development-index-escosura"]
        e_full = e.reindex(range(e.index.min(), e.index.max() + 1)).interpolate()
        d["hdi_source"] = np.where(d.year.isin(u.index), "UNDP", "Escosura-interp")
        d["hdi"] = np.where(d.year.isin(u.index), d.year.map(u), d.year.map(e_full))
        g = gdp[gdp.entity == country].set_index("year")["gdp-per-capita-maddison-project-database"]
        d["gdp_pc"] = d.year.map(g)
        d.insert(0, "country", country)
        panels.append(d)
    p = pd.concat(panels, ignore_index=True).rename(
        columns={"electoral-democracy-index": "polyarchy", "liberal-democracy-index": "libdem"}
    )

    # ---- index maths (identical to the workbook formulas) ---------------------------------------
    a = np.where(p.hdi_source == "UNDP", reg["modern"][0], reg["historical"][0])
    b = np.where(p.hdi_source == "UNDP", reg["modern"][1], reg["historical"][1])
    p["libdem_pred"] = a + b * p.hdi
    p["resid"] = p.libdem - p.libdem_pred
    p["hdi_adj"] = (0.5 + p.resid).clip(0, 1)
    p["index"] = 100 * (
        WEIGHTS["polyarchy"] * p.polyarchy + WEIGHTS["libdem"] * p.libdem
        + WEIGHTS["hdi"] * p.hdi + WEIGHTS["hdi_adj"] * p.hdi_adj
    )
    p = p.sort_values(["country", "year"]).reset_index(drop=True)
    grp = p.groupby("country")
    p["index_chg_3y"] = p["index"] - grp["index"].shift(3)
    p["gdp_yoy"] = grp["gdp_pc"].pct_change(fill_method=None)
    p["strain"] = (p.gdp_yoy <= STRAIN_THRESHOLD).astype(int).where(p.gdp_yoy.notna())
    p.to_csv(PROC / "index_panel.csv", index=False)

    summary = hypothesis_check(p)
    write_workbook(p, reg, summary)
    plot(p)
    print("regressions (a, b, R2, n):", {k: tuple(round(float(x), 3) for x in v) for k, v in reg.items()})
    print(summary.to_string(index=False))


def hypothesis_check(p):
    """Does the index fall faster in the 3 years after a strain year than after a non-strain year?"""
    rows = []
    for country, d in p.groupby("country"):
        d = d.set_index("year")
        fwd = d["index"].shift(-3) - d["index"]  # change over the next 3 years
        s = d["strain"]
        rows.append({
            "country": country,
            "strain_years": ", ".join(str(y) for y in s.index[s == 1]),
            "fwd3y_change_after_strain": round(fwd[s == 1].mean(), 2),
            "fwd3y_change_after_non_strain": round(fwd[s == 0].mean(), 2),
            "n_strain": int((s == 1).sum()),
            "index_min": round(d["index"].min(), 1),
            "index_min_year": int(d["index"].idxmin()),
        })
    return pd.DataFrame(rows)


def write_workbook(p, reg, summary):
    wb = Workbook()
    bold = Font(bold=True)
    edit = PatternFill("solid", fgColor="FFF2CC")

    ws = wb.active
    ws.title = "Weights"
    ws.append(["Component", "Weight (edit yellow cells)", "Meaning"])
    meanings = {
        "polyarchy": "V-Dem electoral democracy index (0-1)",
        "libdem": "V-Dem liberal democracy index (0-1): rights + constraints on the executive",
        "hdi": "Human Development Index level (0-1)",
        "hdi_adj": "0.5 + (actual libdem - libdem predicted from HDI), clipped 0-1",
    }
    for k, v in WEIGHTS.items():
        ws.append([k, v, meanings[k]])
    ws.append(["check: sum should be 1", "=SUM(B2:B5)"])
    ws.append([])
    ws.append(["strain threshold (GDP pc yoy)", STRAIN_THRESHOLD, "yoy change at/below this flags a 'strain year'"])
    for r in range(2, 6):
        ws.cell(r, 2).fill = edit
    ws["B8"].fill = edit
    for c in ws[1]:
        c.font = bold
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 75

    wr = wb.create_sheet("Regression")
    wr.append(["Model: libdem = intercept + slope * HDI (pooled OLS, all country-years)", "intercept", "slope", "R^2", "n"])
    for name, key in [("historical (Escosura HDI, 1870-2020)", "historical"), ("modern (UNDP HDI, 1990-2023)", "modern")]:
        a, b, r2, n = reg[key]
        wr.append([name, a, b, r2, n])
    for c in wr[1]:
        c.font = bold
    wr.column_dimensions["A"].width = 68

    cols = ["year", "polyarchy", "libdem", "hdi", "hdi_source", "libdem_pred", "resid", "hdi_adj", "INDEX",
            "index_chg_3y", "gdp_pc", "gdp_yoy", "strain"]
    for country, d in p.groupby("country"):
        s = wb.create_sheet(country)
        s.append(cols)
        for c in s[1]:
            c.font = bold
        for i, (_, r) in enumerate(d.iterrows(), start=2):
            reg_row = 3 if r.hdi_source == "UNDP" else 2  # Regression sheet row
            nn = lambda v: None if pd.isna(v) else float(v)  # noqa: E731
            s.cell(i, 1, int(r.year))
            s.cell(i, 2, nn(r.polyarchy))
            s.cell(i, 3, nn(r.libdem))
            s.cell(i, 4, nn(r.hdi))
            s.cell(i, 5, r.hdi_source)
            s.cell(i, 6, f"=Regression!$B${reg_row}+Regression!$C${reg_row}*D{i}")
            s.cell(i, 7, f"=C{i}-F{i}")
            s.cell(i, 8, f"=MAX(0,MIN(1,0.5+G{i}))")
            s.cell(i, 9, f"=100*(Weights!$B$2*B{i}+Weights!$B$3*C{i}+Weights!$B$4*D{i}+Weights!$B$5*H{i})")
            if i >= 5:
                s.cell(i, 10, f"=I{i}-I{i-3}")
            s.cell(i, 11, nn(r.gdp_pc))
            if i >= 3 and not pd.isna(r.gdp_yoy):
                s.cell(i, 12, f"=K{i}/K{i-1}-1")
                s.cell(i, 13, f"=IF(L{i}<=Weights!$B$8,1,0)")
        for j in range(1, len(cols) + 1):
            s.column_dimensions[get_column_letter(j)].width = 14
        s.freeze_panes = "B2"
        n = len(d) + 1
        ch = LineChart()
        ch.title = f"{country}: index vs V-Dem components"
        ch.y_axis.title = "0-1 scale / index 0-100"
        ch.height, ch.width = 9, 22
        for col in (2, 3, 9):
            ch.add_data(Reference(s, min_col=col, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=n))
        s.add_chart(ch, "O2")

    sm = wb.create_sheet("Summary", 1)
    sm.append(list(summary.columns))
    for c in sm[1]:
        c.font = bold
    for row in summary.itertuples(index=False):
        sm.append(list(row))
    sm.append([])
    sm.append(["Static values computed in python at default weights; per-country sheets recalculate live when weights change."])
    for j in range(1, len(summary.columns) + 1):
        sm.column_dimensions[get_column_letter(j)].width = 30
    wb.save(OUT / "democratic_resilience_index.xlsx")


def plot(p):
    fig, axes = plt.subplots(len(CASES), 1, figsize=(10, 4 * len(CASES)))
    for ax, (country, (_, _, marker, label)) in zip(axes, CASES.items()):
        d = p[p.country == country]
        single = plt.figure(figsize=(10, 4.2))
        for target in (ax, single.gca()):
            target.plot(d.year, d["index"], color="black", lw=2.2, label="Composite index (0-100)")
            target.plot(d.year, d.polyarchy * 100, color="tab:blue", lw=1.2, ls="--", label="V-Dem electoral x100")
            target.plot(d.year, d.libdem * 100, color="tab:green", lw=1.2, ls="--", label="V-Dem liberal x100")
            target.plot(d.year, d.hdi * 100, color="tab:purple", lw=1.2, ls=":", label="HDI x100")
            for y in d.year[d.strain == 1]:
                target.axvspan(y - 0.5, y + 0.5, color="tab:red", alpha=0.18, lw=0)
            target.axvline(marker, color="grey", ls="-.", lw=1)
            target.set_title(f"{label}  |  red = GDP/capita fell >=2% yoy", fontsize=10)
            target.set_ylabel("score")
            target.set_ylim(0, 100)
            target.xaxis.set_major_locator(MaxNLocator(integer=True))
        ax.legend(fontsize=7, loc="lower left")
        single.gca().legend(fontsize=7, loc="lower left")
        single.tight_layout()
        single.savefig(OUT / f"index_{country.lower()}.png", dpi=130)
        plt.close(single)
    fig.tight_layout()
    fig.savefig(OUT / "index_all_cases.png", dpi=130)
    plt.close(fig)


if __name__ == "__main__":
    main()
